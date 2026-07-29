import csv
import hashlib
import io
import json
import logging
import os
import secrets
import time
import uuid
from collections import Counter, defaultdict, deque
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from pathlib import Path

import jwt
import qrcode
from dotenv import load_dotenv
from flask import Flask, Response, g, jsonify, request, send_file, send_from_directory
from flask_cors import CORS
from sqlalchemy import inspect, or_, text
from sqlalchemy.exc import SQLAlchemyError
from werkzeug.utils import secure_filename

from models import (
    AuditLog,
    Book,
    BorrowRecord,
    Notification,
    PasswordResetToken,
    Reservation,
    User,
    db,
    parse_iso_date,
)
from seed_data import SAMPLE_BOOKS

load_dotenv()

ALLOWED_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif"}
LOGIN_WINDOWS = defaultdict(deque)


class JsonLogFormatter(logging.Formatter):
    def format(self, record):
        payload = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "level": record.levelname,
            "service": os.getenv("SERVICE_NAME", "library-flask-api"),
            "requestId": getattr(record, "request_id", "system"),
            "message": record.getMessage(),
        }
        if hasattr(record, "extra_data"):
            payload.update(record.extra_data)
        if record.exc_info:
            payload["exception"] = self.formatException(record.exc_info)
        return json.dumps(payload, ensure_ascii=False)


def configure_logging():
    formatter = JsonLogFormatter()
    handlers = []
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    handlers.append(stream_handler)
    log_file = (os.getenv("APP_LOG_FILE") or "").strip()
    if log_file:
        os.makedirs(os.path.dirname(log_file), exist_ok=True)
        file_handler = logging.FileHandler(log_file, encoding="utf-8")
        file_handler.setFormatter(formatter)
        handlers.append(file_handler)
    root_logger = logging.getLogger()
    root_logger.handlers.clear()
    for handler in handlers:
        root_logger.addHandler(handler)
    root_logger.setLevel(os.getenv("LOG_LEVEL", "INFO"))


def create_app(test_config=None):
    configure_logging()
    app = Flask(__name__)
    allowed_origins = [item.strip() for item in os.getenv("CORS_ORIGINS", "*").split(",") if item.strip()]
    CORS(
        app,
        resources={r"/*": {"origins": allowed_origins or "*"}},
        allow_headers=["Content-Type", "Authorization", "X-Request-ID"],
        expose_headers=["X-Request-ID", "Content-Disposition"],
    )

    upload_folder = os.getenv("UPLOAD_FOLDER", str(Path(__file__).resolve().parent / "uploads"))
    app.config.update(
        SQLALCHEMY_DATABASE_URI=os.getenv(
            "DATABASE_URL", "postgresql://library_user:library_password@localhost:5432/library_db"
        ),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        JSON_SORT_KEYS=False,
        JWT_SECRET_KEY=os.getenv("JWT_SECRET_KEY", "dev-secret-key-change-me"),
        JWT_EXPIRES_HOURS=int(os.getenv("JWT_EXPIRES_HOURS", "24")),
        MAX_CONTENT_LENGTH=int(os.getenv("MAX_UPLOAD_MB", "8")) * 1024 * 1024,
        UPLOAD_FOLDER=upload_folder,
        LOGIN_LIMIT=int(os.getenv("LOGIN_LIMIT", "8")),
        LOGIN_WINDOW_SECONDS=int(os.getenv("LOGIN_WINDOW_SECONDS", "300")),
    )
    if test_config:
        app.config.update(test_config)
    Path(app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)

    db.init_app(app)
    logger = logging.getLogger(__name__)

    try:
        if os.getenv("ENABLE_XRAY", "false").lower() == "true":
            from aws_xray_sdk.core import patch_all
            patch_all()
            logger.info("AWS X-Ray instrumentation enabled")
    except Exception as error:  # pragma: no cover - optional integration
        logger.warning("AWS X-Ray could not be enabled: %s", error)

    @app.before_request
    def begin_request():
        g.request_id = request.headers.get("X-Request-ID") or str(uuid.uuid4())
        g.request_started_at = time.perf_counter()

    @app.after_request
    def finish_request(response):
        started_at = getattr(g, "request_started_at", None)
        duration_ms = round((time.perf_counter() - started_at) * 1000, 2) if started_at else 0.0
        request_id = getattr(g, "request_id", "unknown")
        response.headers["X-Request-ID"] = request_id
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Permissions-Policy"] = "camera=(self), microphone=(), geolocation=()"
        if request.path != "/health":
            logger.info(
                "HTTP request completed",
                extra={
                    "request_id": request_id,
                    "extra_data": {
                        "event": "http_request",
                        "method": request.method,
                        "path": request.path,
                        "status_code": response.status_code,
                        "duration_ms": duration_ms,
                    },
                },
            )
        return response

    def log_info(message, **kwargs):
        logger.info(message, extra={"request_id": getattr(g, "request_id", "system"), "extra_data": kwargs})

    def log_error(message, **kwargs):
        logger.error(
            message,
            extra={"request_id": getattr(g, "request_id", "system"), "extra_data": kwargs},
            exc_info=True,
        )

    def generate_token(user):
        now = datetime.now(timezone.utc)
        payload = {
            "sub": str(user.id),
            "user_id": user.id,
            "username": user.username,
            "role": user.role,
            "iat": now,
            "exp": now + timedelta(hours=app.config["JWT_EXPIRES_HOURS"]),
        }
        return jwt.encode(payload, app.config["JWT_SECRET_KEY"], algorithm="HS256")

    def get_bearer_token():
        auth_header = request.headers.get("Authorization", "")
        return auth_header.replace("Bearer ", "", 1).strip() if auth_header.startswith("Bearer ") else None

    def token_required(handler):
        @wraps(handler)
        def wrapper(*args, **kwargs):
            token = get_bearer_token()
            if not token:
                return jsonify({"message": "Authorization token is required"}), 401
            try:
                decoded = jwt.decode(token, app.config["JWT_SECRET_KEY"], algorithms=["HS256"])
                user = db.session.get(User, decoded.get("user_id"))
                if not user or not user.is_active:
                    return jsonify({"message": "Tài khoản không tồn tại hoặc đã bị khóa"}), 401
                g.current_user = user
                return handler(*args, **kwargs)
            except jwt.ExpiredSignatureError:
                return jsonify({"message": "Token has expired"}), 401
            except jwt.InvalidTokenError:
                return jsonify({"message": "Invalid token"}), 401
        return wrapper

    def admin_required(handler):
        @wraps(handler)
        @token_required
        def wrapper(*args, **kwargs):
            if g.current_user.role != "admin":
                return jsonify({"message": "Admin permission is required"}), 403
            return handler(*args, **kwargs)
        return wrapper

    def parse_positive_int(value, field_name, minimum=0):
        try:
            parsed = int(value)
        except (TypeError, ValueError):
            raise ValueError(f"{field_name} must be an integer")
        if parsed < minimum:
            raise ValueError(f"{field_name} must be at least {minimum}")
        return parsed

    def due_date_or_default(value):
        due = parse_iso_date(value)
        if value and not due:
            raise ValueError("expected return date must use YYYY-MM-DD")
        due = due or (date.today() + timedelta(days=14))
        if due <= date.today():
            raise ValueError("expected return date must be after today")
        if due > date.today() + timedelta(days=60):
            raise ValueError("expected return date cannot exceed 60 days")
        return due

    def audit(action, entity_type=None, entity_id=None, details=None, user=None):
        actor = user or getattr(g, "current_user", None)
        entry = AuditLog(
            user_id=actor.id if actor else None,
            username=actor.username if actor else None,
            action=action,
            entity_type=entity_type,
            entity_id=str(entity_id) if entity_id is not None else None,
            details=json.dumps(details or {}, ensure_ascii=False, default=str),
            ip_address=request.headers.get("X-Forwarded-For", request.remote_addr),
            request_id=getattr(g, "request_id", None),
        )
        db.session.add(entry)

    def notify(user_id, title, message, notification_type="info", link=None):
        if not user_id:
            return
        db.session.add(
            Notification(
                user_id=int(user_id),
                title=title,
                message=message,
                type=notification_type,
                link=link,
            )
        )

    def send_email(subject, body, recipient):
        if not recipient or not os.getenv("SES_FROM_EMAIL"):
            return False
        try:
            import boto3
            boto3.client("ses", region_name=os.getenv("AWS_REGION", "ap-southeast-2")).send_email(
                Source=os.getenv("SES_FROM_EMAIL"),
                Destination={"ToAddresses": [recipient]},
                Message={
                    "Subject": {"Data": subject, "Charset": "UTF-8"},
                    "Body": {"Text": {"Data": body, "Charset": "UTF-8"}},
                },
            )
            return True
        except Exception as error:  # pragma: no cover - external integration
            logger.warning("SES email failed: %s", error)
            return False

    def check_login_rate_limit(key):
        now = time.time()
        window = app.config["LOGIN_WINDOW_SECONDS"]
        queue = LOGIN_WINDOWS[key]
        while queue and queue[0] < now - window:
            queue.popleft()
        if len(queue) >= app.config["LOGIN_LIMIT"]:
            return False
        queue.append(now)
        return True

    def ensure_schema():
        inspector = inspect(db.engine)

        def columns(table_name):
            if not inspector.has_table(table_name):
                return set()
            return {column["name"] for column in inspector.get_columns(table_name)}

        def add(table_name, column_name, ddl):
            if column_name not in columns(table_name):
                db.session.execute(text(f"ALTER TABLE {table_name} ADD COLUMN {ddl}"))
                db.session.commit()
                logger.info(
                    "Added missing database column",
                    extra={"request_id": "system", "extra_data": {"table": table_name, "column": column_name}},
                )

        user_columns = {
            "full_name": "full_name VARCHAR(255)",
            "email": "email VARCHAR(255)",
            "phone": "phone VARCHAR(20)",
            "cccd": "cccd VARCHAR(20)",
            "avatar": "avatar TEXT",
            "is_active": "is_active BOOLEAN DEFAULT TRUE NOT NULL",
            "failed_login_attempts": "failed_login_attempts INTEGER DEFAULT 0 NOT NULL",
            "locked_until": "locked_until TIMESTAMP WITH TIME ZONE",
        }
        book_columns = {
            "cover": "cover TEXT",
            "isbn": "isbn VARCHAR(32)",
            "publisher": "publisher VARCHAR(180)",
            "publication_year": "publication_year INTEGER",
            "shelf_location": "shelf_location VARCHAR(80)",
            "total_quantity": "total_quantity INTEGER DEFAULT 1 NOT NULL",
            "available_quantity": "available_quantity INTEGER DEFAULT 1 NOT NULL",
            "is_active": "is_active BOOLEAN DEFAULT TRUE NOT NULL",
            "borrow_count": "borrow_count INTEGER DEFAULT 0 NOT NULL",
        }
        record_columns = {
            "customer_name": "customer_name VARCHAR(255)",
            "phone": "phone VARCHAR(20)",
            "cccd": "cccd VARCHAR(20)",
            "expected_return_date": "expected_return_date VARCHAR(50)",
            "deposit": "deposit INTEGER",
            "renew_count": "renew_count INTEGER DEFAULT 0 NOT NULL",
            "last_renewed_at": "last_renewed_at TIMESTAMP WITH TIME ZONE",
            "renewal_status": "renewal_status VARCHAR(30) DEFAULT 'none' NOT NULL",
            "renewal_requested_at": "renewal_requested_at TIMESTAMP WITH TIME ZONE",
            "return_requested_at": "return_requested_at TIMESTAMP WITH TIME ZONE",
            "return_condition": "return_condition VARCHAR(30)",
            "fine_amount": "fine_amount INTEGER DEFAULT 0 NOT NULL",
            "admin_note": "admin_note TEXT",
            "approved_by": "approved_by INTEGER",
        }
        for name, ddl in user_columns.items():
            add("users", name, ddl)
        for name, ddl in book_columns.items():
            add("books", name, ddl)
        for name, ddl in record_columns.items():
            add("borrow_records", name, ddl)
        db.session.execute(text("UPDATE users SET is_active = TRUE WHERE is_active IS NULL"))
        db.session.execute(text("UPDATE users SET failed_login_attempts = 0 WHERE failed_login_attempts IS NULL"))
        db.session.execute(text("UPDATE books SET is_active = TRUE WHERE is_active IS NULL"))
        db.session.execute(text("UPDATE books SET borrow_count = 0 WHERE borrow_count IS NULL"))
        db.session.execute(text("UPDATE borrow_records SET renew_count = 0 WHERE renew_count IS NULL"))
        db.session.execute(text("UPDATE borrow_records SET renewal_status = 'none' WHERE renewal_status IS NULL"))
        db.session.execute(text("UPDATE borrow_records SET fine_amount = 0 WHERE fine_amount IS NULL"))
        db.session.commit()

    def seed_users():
        seeds = [
            (os.getenv("DEFAULT_ADMIN_USERNAME", "admin"), os.getenv("DEFAULT_ADMIN_PASSWORD", "admin123"), "admin", "Quản trị viên thư viện", "admin@cloudlibrary.local"),
            (os.getenv("DEFAULT_USER_USERNAME", "user"), os.getenv("DEFAULT_USER_PASSWORD", "user123"), "user", "Độc giả CloudLibrary", "user@cloudlibrary.local"),
            ("student", "student123", "user", "Nguyễn Minh Anh", "student@cloudlibrary.local"),
        ]
        for username, password, role, full_name, email in seeds:
            user = User.query.filter_by(username=username).first()
            if not user:
                user = User(username=username, role=role, full_name=full_name, email=email, is_active=True)
                user.set_password(password)
                db.session.add(user)
            elif not user.email:
                user.email = email
        db.session.commit()

    def seed_books():
        for item in SAMPLE_BOOKS:
            book = Book.query.filter_by(title=item["title"]).first()
            total = int(item.get("total_quantity", 3))
            if not book:
                book = Book(
                    title=item["title"], author=item["author"], category=item.get("category") or "General",
                    description=item.get("description"), cover=item.get("cover"), isbn=item.get("isbn"),
                    publisher=item.get("publisher"), publication_year=item.get("publication_year"),
                    shelf_location=item.get("shelf_location"), total_quantity=total, available_quantity=total,
                    available=True, is_active=True,
                )
                db.session.add(book)
            else:
                # Always migrate remote demo covers to packaged local files.
                for field in ["cover", "isbn", "publisher", "publication_year", "shelf_location", "description", "category"]:
                    desired = item.get(field)
                    if desired and (field == "cover" or not getattr(book, field, None)):
                        setattr(book, field, desired)
                if (book.total_quantity or 1) <= 1 and not book.borrow_records:
                    book.total_quantity = total
                    book.available_quantity = total
                book.sync_availability()
        db.session.commit()

    with app.app_context():
        db.create_all()
        try:
            ensure_schema()
        except SQLAlchemyError as error:
            db.session.rollback()
            logger.warning("Could not auto-update existing schema: %s", error)
        # New tables are created after old-table columns are migrated.
        db.create_all()
        seed_users()
        seed_books()

    @app.errorhandler(404)
    def not_found(_):
        return jsonify({"message": "Resource not found"}), 404

    @app.errorhandler(413)
    def too_large(_):
        return jsonify({"message": "Ảnh vượt quá dung lượng cho phép"}), 413

    @app.errorhandler(500)
    def internal_error(error):
        log_error("Internal server error", error=str(error))
        db.session.rollback()
        return jsonify({"message": "Internal server error"}), 500

    @app.get("/health")
    def health_check():
        db.session.execute(text("SELECT 1"))
        return jsonify({"status": "ok", "service": "library-flask-api", "time": datetime.now(timezone.utc).isoformat()})

    # Authentication and profile -------------------------------------------------
    @app.post("/auth/register")
    def register():
        payload = request.get_json(silent=True) or {}
        username = (payload.get("username") or "").strip().lower()
        password = payload.get("password") or ""
        email = (payload.get("email") or "").strip().lower() or None
        if not username or len(username) < 3 or not password:
            return jsonify({"message": "Tên đăng nhập phải có ít nhất 3 ký tự"}), 400
        if len(password) < 8:
            return jsonify({"message": "Mật khẩu phải có ít nhất 8 ký tự"}), 400
        if User.query.filter_by(username=username).first():
            return jsonify({"message": "Tên đăng nhập đã tồn tại"}), 409
        if email and User.query.filter_by(email=email).first():
            return jsonify({"message": "Email đã được sử dụng"}), 409
        user = User(
            username=username,
            role="user",
            full_name=(payload.get("fullName") or username).strip(),
            email=email,
            phone=(payload.get("phone") or "").strip() or None,
            cccd=(payload.get("cccd") or "").strip() or None,
            is_active=True,
        )
        user.set_password(password)
        db.session.add(user)
        db.session.flush()
        notify(user.id, "Chào mừng đến CloudLibrary", "Tài khoản độc giả của bạn đã được tạo thành công.", "success", "/dashboard")
        audit("user.register", "user", user.id, {"username": username}, user=user)
        db.session.commit()
        return jsonify({"access_token": generate_token(user), "user": user.to_dict(reveal_sensitive=True)}), 201

    @app.post("/auth/login")
    def login():
        payload = request.get_json(silent=True) or {}
        username = (payload.get("username") or "").strip().lower()
        key = f"{request.remote_addr}:{username}"
        if not check_login_rate_limit(key):
            return jsonify({"message": "Bạn thử đăng nhập quá nhiều lần. Vui lòng chờ vài phút."}), 429
        user = User.query.filter_by(username=username).first()
        if not user or not user.check_password(payload.get("password") or ""):
            if user:
                user.failed_login_attempts = int(user.failed_login_attempts or 0) + 1
                if user.failed_login_attempts >= 5:
                    user.locked_until = datetime.now(timezone.utc) + timedelta(minutes=10)
                    user.failed_login_attempts = 0
                db.session.commit()
            return jsonify({"message": "Sai tên đăng nhập hoặc mật khẩu"}), 401
        if not user.is_active:
            return jsonify({"message": "Tài khoản đã bị khóa"}), 403
        if user.is_locked():
            return jsonify({"message": "Tài khoản tạm khóa do đăng nhập sai nhiều lần"}), 423
        user.failed_login_attempts = 0
        user.locked_until = None
        audit("user.login", "user", user.id, {"username": user.username}, user=user)
        db.session.commit()
        return jsonify({"access_token": generate_token(user), "user": user.to_dict(reveal_sensitive=True)})

    @app.get("/auth/me")
    @token_required
    def get_me():
        return jsonify(g.current_user.to_dict(reveal_sensitive=True))

    @app.put("/profile")
    @token_required
    def update_profile():
        payload = request.get_json(silent=True) or {}
        user = g.current_user
        before = user.to_dict(reveal_sensitive=True)
        for payload_key, attr in [("fullName", "full_name"), ("email", "email"), ("phone", "phone"), ("cccd", "cccd"), ("avatar", "avatar")]:
            if payload_key in payload:
                setattr(user, attr, (payload.get(payload_key) or "").strip() or None)
        if payload.get("newPassword"):
            if not user.check_password(payload.get("currentPassword") or ""):
                return jsonify({"message": "Mật khẩu hiện tại không đúng"}), 400
            if len(payload["newPassword"]) < 8:
                return jsonify({"message": "Mật khẩu mới phải có ít nhất 8 ký tự"}), 400
            user.set_password(payload["newPassword"])
        audit("profile.update", "user", user.id, {"before": before, "after": user.to_dict(reveal_sensitive=False)})
        db.session.commit()
        return jsonify(user.to_dict(reveal_sensitive=True))

    @app.post("/auth/forgot-password")
    def forgot_password():
        payload = request.get_json(silent=True) or {}
        identity = (payload.get("identity") or "").strip().lower()
        user = User.query.filter(or_(User.username == identity, User.email == identity)).first()
        response = {"message": "Nếu tài khoản tồn tại, hướng dẫn đặt lại mật khẩu đã được gửi."}
        if not user:
            return jsonify(response)
        raw_token = secrets.token_urlsafe(32)
        token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
        reset = PasswordResetToken(
            user_id=user.id,
            token_hash=token_hash,
            expires_at=datetime.now(timezone.utc) + timedelta(minutes=30),
        )
        db.session.add(reset)
        db.session.commit()
        sent = send_email(
            "Đặt lại mật khẩu CloudLibrary",
            f"Mã đặt lại mật khẩu của bạn: {raw_token}\nMã có hiệu lực trong 30 phút.",
            user.email,
        )
        if app.config.get("TESTING") or os.getenv("EXPOSE_RESET_TOKEN", "false").lower() == "true":
            response["resetToken"] = raw_token
        response["emailSent"] = sent
        return jsonify(response)

    @app.post("/auth/reset-password")
    def reset_password():
        payload = request.get_json(silent=True) or {}
        raw_token = payload.get("token") or ""
        password = payload.get("password") or ""
        if len(password) < 8:
            return jsonify({"message": "Mật khẩu mới phải có ít nhất 8 ký tự"}), 400
        token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
        reset = PasswordResetToken.query.filter_by(token_hash=token_hash, used_at=None).first()
        if not reset:
            return jsonify({"message": "Mã đặt lại mật khẩu không hợp lệ"}), 400
        expires = reset.expires_at
        if expires.tzinfo is None:
            expires = expires.replace(tzinfo=timezone.utc)
        if expires < datetime.now(timezone.utc):
            return jsonify({"message": "Mã đặt lại mật khẩu đã hết hạn"}), 400
        user = db.session.get(User, reset.user_id)
        user.set_password(password)
        reset.used_at = datetime.now(timezone.utc)
        db.session.commit()
        return jsonify({"message": "Đổi mật khẩu thành công"})

    # Uploads --------------------------------------------------------------------
    def upload_to_storage(file):
        extension = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
        if extension not in ALLOWED_IMAGE_EXTENSIONS:
            raise ValueError("Chỉ hỗ trợ PNG, JPG, WEBP hoặc GIF")
        safe_stem = secure_filename(file.filename.rsplit(".", 1)[0])[:60] or "image"
        filename = f"{safe_stem}-{uuid.uuid4().hex[:10]}.{extension}"
        bucket = os.getenv("AWS_IMAGE_BUCKET")
        if bucket:
            import boto3
            key = f"library-images/{filename}"
            file.stream.seek(0)
            boto3.client("s3", region_name=os.getenv("AWS_REGION", "ap-southeast-2")).upload_fileobj(
                file.stream,
                bucket,
                key,
                ExtraArgs={"ContentType": file.mimetype, "CacheControl": "public,max-age=31536000"},
            )
            base_url = (os.getenv("AWS_IMAGE_BASE_URL") or "").strip()
            # With CloudFront/public base URL, serve assets directly. Otherwise
            # keep the S3 bucket private and proxy images through this API.
            return f"{base_url.rstrip('/')}/{key}" if base_url else f"/media/{key}"
        destination = Path(app.config["UPLOAD_FOLDER"]) / filename
        file.save(destination)
        return f"/uploads/{filename}"

    @app.post("/uploads")
    @token_required
    def upload_image():
        file = request.files.get("file")
        if not file or not file.filename:
            return jsonify({"message": "Vui lòng chọn ảnh"}), 400
        try:
            url = upload_to_storage(file)
            audit("image.upload", "file", url, {"name": file.filename})
            db.session.commit()
            return jsonify({"url": url})
        except ValueError as error:
            return jsonify({"message": str(error)}), 400

    @app.get("/uploads/<path:filename>")
    def serve_upload(filename):
        return send_from_directory(app.config["UPLOAD_FOLDER"], filename, max_age=31536000)

    @app.get("/media/<path:key>")
    def serve_private_s3_media(key):
        bucket = (os.getenv("AWS_IMAGE_BUCKET") or "").strip()
        if not bucket or not key.startswith("uploads/"):
            return jsonify({"message": "Image not found"}), 404
        try:
            import boto3
            result = boto3.client("s3", region_name=os.getenv("AWS_REGION", "ap-southeast-2")).get_object(Bucket=bucket, Key=key)
            payload = io.BytesIO(result["Body"].read())
            return send_file(payload, mimetype=result.get("ContentType") or "application/octet-stream", max_age=31536000)
        except Exception:
            return jsonify({"message": "Image not found"}), 404

    # Users ---------------------------------------------------------------------
    @app.get("/users")
    @admin_required
    def get_users():
        return jsonify([user.to_dict(reveal_sensitive=True) for user in User.query.order_by(User.id).all()])

    @app.put("/users/<int:user_id>/status")
    @admin_required
    def update_user_status(user_id):
        user = db.session.get(User, user_id)
        if not user:
            return jsonify({"message": "User not found"}), 404
        if user.id == g.current_user.id:
            return jsonify({"message": "Không thể khóa chính tài khoản đang đăng nhập"}), 409
        payload = request.get_json(silent=True) or {}
        user.is_active = bool(payload.get("isActive"))
        audit("user.status", "user", user.id, {"isActive": user.is_active})
        db.session.commit()
        return jsonify(user.to_dict(reveal_sensitive=True))

    @app.post("/users/<int:user_id>/reset-password")
    @admin_required
    def admin_reset_password(user_id):
        user = db.session.get(User, user_id)
        if not user:
            return jsonify({"message": "User not found"}), 404
        payload = request.get_json(silent=True) or {}
        password = payload.get("password") or ""
        if len(password) < 8:
            return jsonify({"message": "Mật khẩu phải có ít nhất 8 ký tự"}), 400
        user.set_password(password)
        audit("user.password_reset", "user", user.id)
        db.session.commit()
        return jsonify({"message": "Đặt lại mật khẩu thành công"})

    # Books ---------------------------------------------------------------------
    def book_payload(payload, existing=None):
        title = (payload.get("title") or (existing.title if existing else "")).strip()
        author = (payload.get("author") or (existing.author if existing else "")).strip()
        if not title or not author:
            raise ValueError("title and author are required")
        total = parse_positive_int(payload.get("totalQuantity", existing.total_quantity if existing else 1), "totalQuantity", 1)
        borrowed = existing.borrowed_quantity if existing else 0
        if total < borrowed:
            raise ValueError(f"Tổng số bản không thể thấp hơn {borrowed} bản đang mượn")
        return {
            "title": title,
            "author": author,
            "category": (payload.get("category") or (existing.category if existing else "General")).strip() or "General",
            "description": (payload.get("description") if "description" in payload else (existing.description if existing else None)),
            "cover": (payload.get("cover") if "cover" in payload else (existing.cover if existing else None)),
            "isbn": (payload.get("isbn") if "isbn" in payload else (existing.isbn if existing else None)),
            "publisher": (payload.get("publisher") if "publisher" in payload else (existing.publisher if existing else None)),
            "publication_year": int(payload["publicationYear"]) if payload.get("publicationYear") else None,
            "shelf_location": (payload.get("shelfLocation") if "shelfLocation" in payload else (existing.shelf_location if existing else None)),
            "total_quantity": total,
            "available_quantity": total - borrowed,
        }

    @app.get("/books")
    @token_required
    def get_books():
        query = Book.query.filter_by(is_active=True)
        keyword = (request.args.get("q") or "").strip()
        if keyword:
            term = f"%{keyword}%"
            query = query.filter(or_(Book.title.ilike(term), Book.author.ilike(term), Book.category.ilike(term), Book.isbn.ilike(term)))
        return jsonify([book.to_dict() for book in query.order_by(Book.id).all()])

    @app.get("/admin/books")
    @admin_required
    def get_admin_books():
        return jsonify([book.to_dict() for book in Book.query.order_by(Book.id).all()])

    @app.get("/books/<int:book_id>")
    @token_required
    def get_book(book_id):
        book = db.session.get(Book, book_id)
        if not book:
            return jsonify({"message": "Book not found"}), 404
        data = book.to_dict(include_records=g.current_user.role == "admin")
        if g.current_user.role != "admin":
            data["myRecords"] = [r.to_dict() for r in book.borrow_records if r.user_id == str(g.current_user.id)]
        return jsonify(data)

    @app.get("/books/<int:book_id>/qr")
    def book_qr(book_id):
        target = f"{request.host_url.rstrip('/')}/?book={book_id}"
        image = qrcode.make(target)
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return Response(buffer.getvalue(), mimetype="image/png")

    @app.post("/books")
    @admin_required
    def create_book():
        payload = request.get_json(silent=True) or {}
        try:
            values = book_payload(payload)
            book = Book(**values, available=True, is_active=True)
            book.sync_availability()
            db.session.add(book)
            db.session.flush()
            audit("book.create", "book", book.id, book.to_dict())
            db.session.commit()
            return jsonify(book.to_dict()), 201
        except ValueError as error:
            return jsonify({"message": str(error)}), 400

    @app.put("/books/<int:book_id>")
    @admin_required
    def update_book(book_id):
        book = db.session.get(Book, book_id)
        if not book:
            return jsonify({"message": "Book not found"}), 404
        payload = request.get_json(silent=True) or {}
        try:
            before = book.to_dict()
            for attr, value in book_payload(payload, book).items():
                setattr(book, attr, value)
            book.sync_availability()
            audit("book.update", "book", book.id, {"before": before, "after": book.to_dict()})
            db.session.commit()
            return jsonify(book.to_dict())
        except ValueError as error:
            return jsonify({"message": str(error)}), 400

    @app.delete("/books/<int:book_id>")
    @admin_required
    def archive_book(book_id):
        book = db.session.get(Book, book_id)
        if not book:
            return jsonify({"message": "Book not found"}), 404
        if book.borrowed_quantity > 0:
            return jsonify({"message": "Không thể lưu trữ sách đang được mượn"}), 409
        book.is_active = False
        book.sync_availability()
        audit("book.archive", "book", book.id)
        db.session.commit()
        return jsonify(book.to_dict())

    @app.post("/books/<int:book_id>/restore")
    @admin_required
    def restore_book(book_id):
        book = db.session.get(Book, book_id)
        if not book:
            return jsonify({"message": "Book not found"}), 404
        book.is_active = True
        book.sync_availability()
        audit("book.restore", "book", book.id)
        db.session.commit()
        return jsonify(book.to_dict())

    # Reservations ---------------------------------------------------------------
    def reservation_rows(query):
        rows = query.order_by(Reservation.book_id, Reservation.created_at).all()
        counters = defaultdict(int)
        result = []
        for row in rows:
            if row.status in {"waiting", "notified"}:
                counters[row.book_id] += 1
                position = counters[row.book_id]
            else:
                position = None
            result.append(row.to_dict(position))
        return result

    @app.post("/books/<int:book_id>/reservations")
    @token_required
    def reserve_book(book_id):
        book = db.session.get(Book, book_id)
        if not book or not book.is_active:
            return jsonify({"message": "Book not found"}), 404
        existing = Reservation.query.filter_by(book_id=book_id, user_id=g.current_user.id).filter(Reservation.status.in_(["waiting", "notified"])).first()
        if existing:
            return jsonify({"message": "Bạn đã có trong hàng chờ của sách này"}), 409
        active_loan = BorrowRecord.query.filter_by(book_id=book_id, user_id=str(g.current_user.id)).filter(BorrowRecord.status.in_(["active", "return_requested"])).first()
        if active_loan:
            return jsonify({"message": "Bạn đang mượn sách này"}), 409
        reservation = Reservation(book_id=book_id, user_id=g.current_user.id, status="waiting")
        db.session.add(reservation)
        db.session.flush()
        notify(g.current_user.id, "Đã đặt trước sách", f"Bạn đã vào hàng chờ cho “{book.title}”.", "info", "/reservations")
        audit("reservation.create", "reservation", reservation.id, {"bookId": book_id})
        db.session.commit()
        position = Reservation.query.filter_by(book_id=book_id, status="waiting").filter(Reservation.created_at <= reservation.created_at).count()
        return jsonify(reservation.to_dict(position)), 201

    @app.get("/reservations")
    @token_required
    def get_reservations():
        query = Reservation.query
        if g.current_user.role != "admin":
            query = query.filter_by(user_id=g.current_user.id)
        return jsonify(reservation_rows(query))

    @app.delete("/reservations/<int:reservation_id>")
    @token_required
    def cancel_reservation(reservation_id):
        reservation = db.session.get(Reservation, reservation_id)
        if not reservation:
            return jsonify({"message": "Reservation not found"}), 404
        if g.current_user.role != "admin" and reservation.user_id != g.current_user.id:
            return jsonify({"message": "Permission denied"}), 403
        if reservation.status not in {"waiting", "notified"}:
            return jsonify({"message": "Reservation cannot be cancelled"}), 409
        reservation.status = "cancelled"
        audit("reservation.cancel", "reservation", reservation.id)
        db.session.commit()
        return jsonify(reservation.to_dict())

    @app.post("/reservations/<int:reservation_id>/notify")
    @admin_required
    def notify_reservation(reservation_id):
        reservation = db.session.get(Reservation, reservation_id)
        if not reservation:
            return jsonify({"message": "Reservation not found"}), 404
        reservation.status = "notified"
        reservation.notified_at = datetime.now(timezone.utc)
        reservation.expires_at = datetime.now(timezone.utc) + timedelta(days=2)
        notify(reservation.user_id, "Sách đặt trước đã có", f"“{reservation.book.title}” đã sẵn sàng. Vui lòng mượn trong 48 giờ.", "success", "/reservations")
        if reservation.user and reservation.user.email:
            send_email("Sách đặt trước đã có", f"Sách {reservation.book.title} đã sẵn sàng trong 48 giờ.", reservation.user.email)
        audit("reservation.notify", "reservation", reservation.id)
        db.session.commit()
        return jsonify(reservation.to_dict(1))

    # Borrowing ------------------------------------------------------------------
    def resolve_borrower(payload):
        if g.current_user.role != "admin":
            return g.current_user
        user_id = payload.get("user_id") or payload.get("userId")
        return db.session.get(User, int(user_id)) if user_id else None

    @app.post("/borrow")
    @token_required
    def borrow_book():
        payload = request.get_json(silent=True) or {}
        try:
            book_id = int(payload.get("book_id") or payload.get("bookId"))
            borrower = resolve_borrower(payload)
            if not borrower or not borrower.is_active:
                return jsonify({"message": "Vui lòng chọn độc giả hợp lệ"}), 400
            book = db.session.get(Book, book_id)
            if not book or not book.is_active:
                return jsonify({"message": "Book not found"}), 404
            if book.available_quantity <= 0:
                return jsonify({"message": "Sách đã hết. Hãy sử dụng chức năng đặt trước."}), 409
            existing = BorrowRecord.query.filter_by(book_id=book.id, user_id=str(borrower.id)).filter(BorrowRecord.status.in_(["active", "return_requested"])).first()
            if existing:
                return jsonify({"message": "Độc giả đang mượn sách này"}), 409
            first_reservation = Reservation.query.filter_by(book_id=book.id).filter(Reservation.status.in_(["waiting", "notified"])).order_by(Reservation.created_at).first()
            if first_reservation and first_reservation.user_id != borrower.id:
                return jsonify({"message": "Bản sách đang được ưu tiên cho độc giả đứng đầu hàng chờ"}), 409
            due = due_date_or_default(payload.get("returnDate") or payload.get("expectedReturnDate"))
            deposit = parse_positive_int(payload.get("deposit", 150000), "deposit", 0)
            book.available_quantity -= 1
            book.borrow_count = int(book.borrow_count or 0) + 1
            book.sync_availability()
            record = BorrowRecord(
                book_id=book.id,
                user_id=str(borrower.id),
                customer_name=borrower.full_name or borrower.username,
                phone=borrower.phone,
                cccd=borrower.cccd,
                expected_return_date=due.isoformat(),
                deposit=deposit,
                status="active",
                renewal_status="none",
            )
            db.session.add(record)
            if first_reservation and first_reservation.user_id == borrower.id:
                first_reservation.status = "fulfilled"
                first_reservation.fulfilled_at = datetime.now(timezone.utc)
            db.session.flush()
            notify(borrower.id, "Mượn sách thành công", f"Bạn đã mượn “{book.title}”, hạn trả {due.isoformat()}.", "success", "/my-books")
            audit("borrow.create", "borrow_record", record.id, {"bookId": book.id, "userId": borrower.id})
            db.session.commit()
            return jsonify({"message": "Borrow book successfully", "book": book.to_dict(), "record": record.to_dict(reveal_sensitive=g.current_user.role == "admin")})
        except (TypeError, ValueError) as error:
            return jsonify({"message": str(error)}), 400
        except SQLAlchemyError as error:
            db.session.rollback()
            log_error("Borrow request failed", event="borrow_failed", error=str(error))
            return jsonify({"message": "Borrow request failed"}), 500

    def get_record_or_404(record_id):
        return db.session.get(BorrowRecord, record_id)

    @app.post("/borrow-records/<int:record_id>/request-return")
    @token_required
    def request_return(record_id):
        record = get_record_or_404(record_id)
        if not record:
            return jsonify({"message": "Borrow record not found"}), 404
        if g.current_user.role != "admin" and record.user_id != str(g.current_user.id):
            return jsonify({"message": "Permission denied"}), 403
        if record.status != "active":
            return jsonify({"message": "Phiếu không ở trạng thái đang mượn"}), 409
        record.status = "return_requested"
        record.return_requested_at = datetime.now(timezone.utc)
        admins = User.query.filter_by(role="admin", is_active=True).all()
        for admin in admins:
            notify(admin.id, "Yêu cầu trả sách", f"{record.customer_name} yêu cầu trả “{record.book.title}”.", "warning", "/loans")
        audit("return.request", "borrow_record", record.id)
        db.session.commit()
        return jsonify(record.to_dict())

    @app.post("/borrow-records/<int:record_id>/approve-return")
    @admin_required
    def approve_return(record_id):
        record = get_record_or_404(record_id)
        if not record:
            return jsonify({"message": "Borrow record not found"}), 404
        if record.status not in {"active", "return_requested"}:
            return jsonify({"message": "Phiếu đã được xử lý"}), 409
        payload = request.get_json(silent=True) or {}
        condition = payload.get("condition") or "good"
        if condition not in {"good", "minor_damage", "major_damage", "lost"}:
            return jsonify({"message": "Tình trạng sách không hợp lệ"}), 400
        fine = parse_positive_int(payload.get("fineAmount", 0), "fineAmount", 0)
        record.return_condition = condition
        record.fine_amount = fine
        record.admin_note = (payload.get("note") or "").strip() or None
        record.approved_by = g.current_user.id
        record.return_date = datetime.now(timezone.utc)
        if condition == "lost":
            record.status = "lost"
            record.book.total_quantity = max(record.book.borrowed_quantity, record.book.total_quantity - 1)
        else:
            record.status = "returned"
            record.book.available_quantity = min(record.book.total_quantity, record.book.available_quantity + 1)
        record.book.sync_availability()
        notify(int(record.user_id), "Đã xác nhận trả sách", f"Thủ thư đã xác nhận trả “{record.book.title}”. Phí phát sinh: {fine:,} VNĐ.", "success", "/my-books")
        audit("return.approve", "borrow_record", record.id, {"condition": condition, "fineAmount": fine})
        db.session.commit()
        return jsonify(record.to_dict(reveal_sensitive=True))

    @app.post("/return")
    @admin_required
    def legacy_return():
        payload = request.get_json(silent=True) or {}
        record_id = payload.get("record_id") or payload.get("recordId")
        if not record_id:
            return jsonify({"message": "record_id is required"}), 400
        return approve_return(int(record_id))

    @app.post("/borrow-records/<int:record_id>/renew")
    @token_required
    def request_or_approve_renew(record_id):
        record = get_record_or_404(record_id)
        if not record:
            return jsonify({"message": "Borrow record not found"}), 404
        if g.current_user.role != "admin" and record.user_id != str(g.current_user.id):
            return jsonify({"message": "Permission denied"}), 403
        if record.status != "active" or record.is_overdue():
            return jsonify({"message": "Phiếu quá hạn hoặc không hoạt động không thể gia hạn"}), 409
        if int(record.renew_count or 0) >= 2:
            return jsonify({"message": "Đã đạt tối đa 2 lần gia hạn"}), 409
        if Reservation.query.filter_by(book_id=record.book_id).filter(Reservation.status.in_(["waiting", "notified"])).first():
            return jsonify({"message": "Không thể gia hạn vì sách đang có người đặt trước"}), 409
        if g.current_user.role == "admin":
            current_due = record.due_date() or date.today()
            record.expected_return_date = (current_due + timedelta(days=7)).isoformat()
            record.renew_count = int(record.renew_count or 0) + 1
            record.last_renewed_at = datetime.now(timezone.utc)
            record.renewal_status = "approved"
            notify(int(record.user_id), "Gia hạn được duyệt", f"Hạn trả “{record.book.title}” đã được cộng thêm 7 ngày.", "success", "/my-books")
            audit("renew.approve", "borrow_record", record.id)
        else:
            if record.renewal_status == "pending":
                return jsonify({"message": "Yêu cầu gia hạn đang chờ duyệt"}), 409
            record.renewal_status = "pending"
            record.renewal_requested_at = datetime.now(timezone.utc)
            for admin in User.query.filter_by(role="admin", is_active=True).all():
                notify(admin.id, "Yêu cầu gia hạn", f"{record.customer_name} yêu cầu gia hạn “{record.book.title}”.", "info", "/loans")
            audit("renew.request", "borrow_record", record.id)
        db.session.commit()
        return jsonify(record.to_dict(reveal_sensitive=g.current_user.role == "admin"))

    @app.post("/borrow-records/<int:record_id>/renew-decision")
    @admin_required
    def renew_decision(record_id):
        record = get_record_or_404(record_id)
        if not record:
            return jsonify({"message": "Borrow record not found"}), 404
        payload = request.get_json(silent=True) or {}
        approved = bool(payload.get("approved"))
        if record.renewal_status != "pending":
            return jsonify({"message": "Không có yêu cầu gia hạn đang chờ"}), 409
        if approved:
            current_due = record.due_date() or date.today()
            record.expected_return_date = (current_due + timedelta(days=7)).isoformat()
            record.renew_count = int(record.renew_count or 0) + 1
            record.last_renewed_at = datetime.now(timezone.utc)
            record.renewal_status = "approved"
            title, message, kind = "Gia hạn được duyệt", f"Hạn trả “{record.book.title}” đã được cộng 7 ngày.", "success"
        else:
            record.renewal_status = "rejected"
            title, message, kind = "Gia hạn bị từ chối", f"Yêu cầu gia hạn “{record.book.title}” không được chấp thuận.", "warning"
        notify(int(record.user_id), title, message, kind, "/my-books")
        audit("renew.decision", "borrow_record", record.id, {"approved": approved})
        db.session.commit()
        return jsonify(record.to_dict(reveal_sensitive=True))

    @app.get("/borrow-records")
    @admin_required
    def get_borrow_records():
        return jsonify([record.to_dict(reveal_sensitive=True) for record in BorrowRecord.query.order_by(BorrowRecord.id.desc()).all()])

    @app.get("/my-borrow-records")
    @token_required
    def get_my_borrow_records():
        records = BorrowRecord.query.filter_by(user_id=str(g.current_user.id)).order_by(BorrowRecord.id.desc()).all()
        return jsonify([record.to_dict(reveal_sensitive=True) for record in records])

    # Notifications and alerts ---------------------------------------------------
    def due_alerts_for(user):
        query = BorrowRecord.query.filter(BorrowRecord.status.in_(["active", "return_requested"]))
        if user.role != "admin":
            query = query.filter_by(user_id=str(user.id))
        alerts = []
        for record in query.order_by(BorrowRecord.expected_return_date.asc()).all():
            days = record.days_remaining()
            if days is None or days > 3:
                continue
            alerts.append({
                "id": record.id,
                "type": "overdue" if days < 0 else "due_soon",
                "severity": "danger" if days < 0 else "warning",
                "message": f"{record.book.title} đã quá hạn {abs(days)} ngày" if days < 0 else f"{record.book.title} còn {days} ngày đến hạn",
                "record": record.to_dict(reveal_sensitive=user.role == "admin"),
            })
        return alerts

    @app.get("/alerts")
    @token_required
    def get_alerts():
        return jsonify(due_alerts_for(g.current_user))

    @app.get("/notifications")
    @token_required
    def get_notifications():
        items = Notification.query.filter_by(user_id=g.current_user.id).order_by(Notification.id.desc()).limit(50).all()
        return jsonify([item.to_dict() for item in items])

    @app.post("/notifications/<int:notification_id>/read")
    @token_required
    def mark_notification(notification_id):
        item = db.session.get(Notification, notification_id)
        if not item or item.user_id != g.current_user.id:
            return jsonify({"message": "Notification not found"}), 404
        item.is_read = True
        db.session.commit()
        return jsonify(item.to_dict())

    @app.post("/notifications/read-all")
    @token_required
    def mark_all_notifications():
        Notification.query.filter_by(user_id=g.current_user.id, is_read=False).update({"is_read": True})
        db.session.commit()
        return jsonify({"message": "Đã đánh dấu tất cả là đã đọc"})

    # Dashboard, search and assistant -------------------------------------------
    @app.get("/dashboard")
    @token_required
    def get_dashboard():
        active_books = Book.query.filter_by(is_active=True).all()
        record_query = BorrowRecord.query
        if g.current_user.role != "admin":
            record_query = record_query.filter_by(user_id=str(g.current_user.id))
        records = record_query.order_by(BorrowRecord.id.desc()).all()
        active_records = [r for r in records if r.status in {"active", "return_requested"}]
        overdue_records = [r for r in active_records if r.is_overdue()]
        category_counts = Counter(book.category or "General" for book in active_books)
        top_books = sorted(active_books, key=lambda item: (item.borrow_count or 0, item.title), reverse=True)[:5]
        total_copies = sum(book.total_quantity or 0 for book in active_books)
        available_copies = sum(book.available_quantity or 0 for book in active_books)
        trend = Counter()
        for record in records:
            if record.borrow_date:
                trend[record.borrow_date.date().isoformat()] += 1
        trend_items = []
        for offset in range(13, -1, -1):
            day = (date.today() - timedelta(days=offset)).isoformat()
            trend_items.append({"date": day, "count": trend.get(day, 0)})
        data = {
            "role": g.current_user.role,
            "totalTitles": len(active_books),
            "totalCopies": total_copies,
            "availableCopies": available_copies,
            "borrowedCopies": max(0, total_copies - available_copies),
            "activeLoans": len(active_records),
            "overdueLoans": len(overdue_records),
            "returnRequests": len([r for r in records if r.status == "return_requested"]),
            "renewalRequests": len([r for r in records if r.renewal_status == "pending"]),
            "returnedLoans": len([r for r in records if r.status == "returned"]),
            "lostLoans": len([r for r in records if r.status == "lost"]),
            "totalFines": sum(int(r.fine_amount or 0) for r in records),
            "totalUsers": User.query.count() if g.current_user.role == "admin" else None,
            "activeReservations": Reservation.query.filter(Reservation.status.in_(["waiting", "notified"])).count() if g.current_user.role == "admin" else Reservation.query.filter_by(user_id=g.current_user.id).filter(Reservation.status.in_(["waiting", "notified"])).count(),
            "categories": [{"name": name, "count": count} for name, count in category_counts.most_common(8)],
            "topBooks": [book.to_dict() for book in top_books],
            "recentRecords": [record.to_dict(reveal_sensitive=g.current_user.role == "admin") for record in records[:6]],
            "borrowTrend": trend_items,
        }
        return jsonify(data)

    @app.get("/recommendations")
    @token_required
    def get_recommendations():
        user_records = BorrowRecord.query.filter_by(user_id=str(g.current_user.id)).all()
        borrowed_book_ids = {record.book_id for record in user_records}
        preferred_categories = Counter(record.book.category for record in user_records if record.book and record.book.category)
        candidates = Book.query.filter_by(is_active=True).filter(Book.available_quantity > 0).all()
        def score(book):
            return preferred_categories.get(book.category, 0) * 100 + (25 if book.id not in borrowed_book_ids else 0) + int(book.borrow_count or 0)
        ranked = sorted(candidates, key=lambda book: (score(book), book.borrow_count or 0, book.title), reverse=True)[:8]
        result = []
        for book in ranked:
            data = book.to_dict()
            if preferred_categories.get(book.category, 0):
                data["recommendationReason"] = f"Phù hợp với thể loại {book.category} bạn đã mượn"
            elif book.borrow_count:
                data["recommendationReason"] = "Được nhiều độc giả lựa chọn"
            else:
                data["recommendationReason"] = "Sách mới trong thư viện"
            result.append(data)
        return jsonify(result)

    @app.get("/search")
    @token_required
    def smart_search():
        query = (request.args.get("q") or "").strip().lower()
        if not query:
            return jsonify([])
        category_aliases = {
            "lập trình": "programming", "python": "python", "web": "web", "aws": "cloud",
            "cơ sở dữ liệu": "database", "ai": "artificial intelligence", "trí tuệ nhân tạo": "artificial intelligence",
            "bảo mật": "security", "mạng": "network", "thuật toán": "algorithm",
        }
        expanded = [query] + [value for key, value in category_aliases.items() if key in query]
        books = Book.query.filter_by(is_active=True).all()
        ranked = []
        for book in books:
            haystack = " ".join(filter(None, [book.title, book.author, book.category, book.description, book.publisher])).lower()
            score = sum(3 if term in book.title.lower() else 1 for term in expanded if term in haystack)
            score += min(3, int(book.borrow_count or 0) // 5)
            if score:
                ranked.append((score, book))
        return jsonify([book.to_dict() for _, book in sorted(ranked, key=lambda item: (-item[0], item[1].title))[:20]])

    @app.post("/assistant/chat")
    @token_required
    def assistant_chat():
        payload = request.get_json(silent=True) or {}
        message = (payload.get("message") or "").strip().lower()
        if not message:
            return jsonify({"message": "Hãy nhập câu hỏi"}), 400
        if "đang mượn" in message or "sách của tôi" in message:
            records = BorrowRecord.query.filter_by(user_id=str(g.current_user.id)).filter(BorrowRecord.status.in_(["active", "return_requested"])).all()
            reply = "Bạn chưa mượn sách nào." if not records else "Bạn đang mượn: " + "; ".join(f"{r.book.title} (hạn {r.expected_return_date})" for r in records[:8])
            return jsonify({"reply": reply, "books": []})
        if "quá hạn" in message or "sắp đến hạn" in message:
            alerts = due_alerts_for(g.current_user)
            reply = "Bạn không có cảnh báo đến hạn." if not alerts else "; ".join(item["message"] for item in alerts[:8])
            return jsonify({"reply": reply, "books": []})
        keywords = message.replace("gợi ý", "").replace("tìm", "").replace("sách", "").strip()
        with app.test_request_context(f"/search?q={keywords}"):
            pass
        books = Book.query.filter_by(is_active=True).all()
        ranked = [b for b in books if keywords and keywords in " ".join(filter(None, [b.title, b.author, b.category, b.description])).lower()]
        if not ranked:
            ranked = sorted([b for b in books if b.available_quantity > 0], key=lambda b: b.borrow_count or 0, reverse=True)[:5]
        return jsonify({"reply": f"Tôi tìm thấy {len(ranked[:5])} gợi ý phù hợp.", "books": [b.to_dict() for b in ranked[:5]]})

    # Audit and reports ----------------------------------------------------------
    @app.get("/audit-logs")
    @admin_required
    def get_audit_logs():
        limit = min(500, max(1, int(request.args.get("limit", 150))))
        return jsonify([item.to_dict() for item in AuditLog.query.order_by(AuditLog.id.desc()).limit(limit).all()])

    def report_rows():
        rows = []
        for record in BorrowRecord.query.order_by(BorrowRecord.id.desc()).all():
            rows.append({
                "Mã phiếu": record.id,
                "Độc giả": record.customer_name,
                "Sách": record.book.title if record.book else "",
                "Ngày mượn": record.borrow_date.strftime("%Y-%m-%d") if record.borrow_date else "",
                "Hạn trả": record.expected_return_date or "",
                "Ngày trả": record.return_date.strftime("%Y-%m-%d") if record.return_date else "",
                "Trạng thái": record.status,
                "Tình trạng sách": record.return_condition or "",
                "Tiền phạt (VNĐ)": int(record.fine_amount or 0),
                "Ghi chú": record.admin_note or "",
            })
        return rows

    @app.post("/books/import")
    @admin_required
    def import_books():
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return jsonify({"message": "Vui lòng chọn tệp CSV hoặc XLSX"}), 400
        extension = upload.filename.rsplit(".", 1)[-1].lower() if "." in upload.filename else ""
        rows = []
        try:
            if extension == "csv":
                content = upload.read().decode("utf-8-sig")
                rows = list(csv.DictReader(io.StringIO(content)))
            elif extension == "xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(upload, read_only=True, data_only=True)
                sheet = workbook.active
                values = list(sheet.iter_rows(values_only=True))
                if values:
                    headers = [str(item or "").strip() for item in values[0]]
                    rows = [dict(zip(headers, row)) for row in values[1:]]
            else:
                return jsonify({"message": "Chỉ hỗ trợ tệp CSV hoặc XLSX"}), 400
        except Exception as error:
            return jsonify({"message": f"Không thể đọc tệp: {error}"}), 400

        aliases = {
            "title": ["title", "Tên sách", "ten_sach"],
            "author": ["author", "Tác giả", "tac_gia"],
            "category": ["category", "Thể loại", "the_loai"],
            "isbn": ["isbn", "ISBN"],
            "publisher": ["publisher", "Nhà xuất bản", "nha_xuat_ban"],
            "publication_year": ["publicationYear", "publication_year", "Năm xuất bản", "nam_xuat_ban"],
            "shelf_location": ["shelfLocation", "shelf_location", "Vị trí kệ", "vi_tri_ke"],
            "total_quantity": ["totalQuantity", "total_quantity", "Tổng số bản", "tong_so_ban"],
            "description": ["description", "Mô tả", "mo_ta"],
        }

        def read_value(row, keys):
            for key in keys:
                value = row.get(key)
                if value is not None and str(value).strip() != "":
                    return str(value).strip()
            return ""

        created = 0
        updated = 0
        errors = []
        for index, row in enumerate(rows, start=2):
            title = read_value(row, aliases["title"])
            author = read_value(row, aliases["author"])
            if not title or not author:
                errors.append(f"Dòng {index}: thiếu tên sách hoặc tác giả")
                continue
            isbn = read_value(row, aliases["isbn"])
            book = Book.query.filter_by(isbn=isbn).first() if isbn else Book.query.filter_by(title=title, author=author).first()
            is_new = book is None
            if is_new:
                book = Book(title=title, author=author, available=True, is_active=True)
                db.session.add(book)
            try:
                total = max(1, int(float(read_value(row, aliases["total_quantity"]) or 1)))
                borrowed = len([record for record in book.borrow_records if record.status == "active"]) if not is_new else 0
                if total < borrowed:
                    raise ValueError(f"tổng số bản phải >= {borrowed}")
                book.title = title
                book.author = author
                book.category = read_value(row, aliases["category"]) or "Khác"
                book.isbn = isbn or book.isbn
                book.publisher = read_value(row, aliases["publisher"]) or None
                year = read_value(row, aliases["publication_year"])
                book.publication_year = int(float(year)) if year else None
                book.shelf_location = read_value(row, aliases["shelf_location"]) or None
                book.description = read_value(row, aliases["description"]) or None
                book.total_quantity = total
                book.available_quantity = max(0, total - borrowed)
                book.is_active = True
                book.sync_availability()
                created += int(is_new)
                updated += int(not is_new)
            except Exception as error:
                if is_new:
                    db.session.expunge(book)
                errors.append(f"Dòng {index}: {error}")
        db.session.commit()
        audit("books_imported", "book", details={"created": created, "updated": updated, "errors": errors[:20]})
        db.session.commit()
        return jsonify({"created": created, "updated": updated, "errors": errors})

    def book_report_rows():
        return [{
            "Mã sách": book.id,
            "Tên sách": book.title,
            "Tác giả": book.author,
            "Thể loại": book.category,
            "ISBN": book.isbn or "",
            "Nhà xuất bản": book.publisher or "",
            "Năm xuất bản": book.publication_year or "",
            "Vị trí kệ": book.shelf_location or "",
            "Tổng số bản": book.total_quantity,
            "Còn lại": book.available_quantity,
            "Đang mượn": book.borrowed_quantity,
            "Trạng thái": "Hoạt động" if book.is_active else "Lưu trữ",
        } for book in Book.query.order_by(Book.id.asc()).all()]

    @app.get("/reports/books")
    @admin_required
    def export_books():
        file_format = (request.args.get("format") or "xlsx").lower()
        rows = book_report_rows()
        fields = list(rows[0].keys()) if rows else ["Mã sách", "Tên sách", "Tác giả"]
        timestamp = datetime.now().strftime("%Y%m%d-%H%M")
        if file_format == "csv":
            buffer = io.StringIO()
            buffer.write("\ufeff")
            writer = csv.DictWriter(buffer, fieldnames=fields)
            writer.writeheader(); writer.writerows(rows)
            return Response(buffer.getvalue(), mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename=book-catalog-{timestamp}.csv"})
        if file_format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font
            workbook = Workbook(); sheet = workbook.active; sheet.title = "Books"
            sheet.append(fields)
            for row in rows: sheet.append([row.get(field) for field in fields])
            for cell in sheet[1]: cell.font = Font(bold=True)
            output = io.BytesIO(); workbook.save(output); output.seek(0)
            return send_file(output, as_attachment=True, download_name=f"book-catalog-{timestamp}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if file_format == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
            output = io.BytesIO(); document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=18, leftMargin=18, topMargin=18, bottomMargin=18)
            data = [fields] + [[str(row.get(field, ""))[:45] for field in fields] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17375E")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTSIZE", (0,0), (-1,-1), 6), ("GRID", (0,0), (-1,-1), .25, colors.grey)]))
            document.build([table]); output.seek(0)
            return send_file(output, as_attachment=True, download_name=f"book-catalog-{timestamp}.pdf", mimetype="application/pdf")
        return jsonify({"message": "Unsupported report format"}), 400

    @app.get("/reports/borrowings")
    @admin_required
    def export_borrowings():
        file_format = (request.args.get("format") or "csv").lower()
        rows = report_rows()
        fields = list(rows[0].keys()) if rows else ["Mã phiếu", "Độc giả", "Sách", "Trạng thái"]
        timestamp = datetime.now().strftime("%Y%m%d-%H%M")
        if file_format == "csv":
            buffer = io.StringIO()
            buffer.write("\ufeff")
            writer = csv.DictWriter(buffer, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
            return Response(buffer.getvalue(), mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename=borrow-report-{timestamp}.csv"})
        if file_format == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Borrowings"
            sheet.append(fields)
            for row in rows:
                sheet.append([row.get(field) for field in fields])
            for cell in sheet[1]:
                cell.font = cell.font.copy(bold=True)
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)
            return send_file(output, as_attachment=True, download_name=f"borrow-report-{timestamp}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if file_format == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
            output = io.BytesIO()
            document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
            data = [fields] + [[str(row.get(field, ""))[:50] for field in fields] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17375E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            document.build([table])
            output.seek(0)
            return send_file(output, as_attachment=True, download_name=f"borrow-report-{timestamp}.pdf", mimetype="application/pdf")
        return jsonify({"message": "Unsupported report format"}), 400

    return app


app = create_app()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)), debug=os.getenv("FLASK_DEBUG", "false").lower() == "true")
